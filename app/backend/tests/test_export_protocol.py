from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

import main
from export_protocol import parse_summary_sections


EXPORT_PAYLOAD = {
    "format": "txt",
    "metadata": {
        "committee": "Hauptausschuss",
        "date": "2026-06-30",
        "location": "Rathaus",
        "title": "Sitzung Hauptausschuss",
        "participants": ["Alice Beispiel", "Bob Beispiel"],
    },
    "appendix": {
        "include_speaker_list": True,
        "include_transcript": True,
        "group_transcript_by_top": False,
        "include_generation_note": True,
    },
    "tops": ["Begrüßung", "Haushalt"],
    "transcript": [
        {"speaker": "SPEAKER_00", "text": "Ich eröffne die Sitzung.", "start": 0, "end": 2},
        {"speaker": "SPEAKER_01", "text": "Der Haushalt wird beraten.", "start": 3, "end": 6},
    ],
    "assignments": [0, 1],
    "speaker_names": {"SPEAKER_00": "Alice", "SPEAKER_01": "Bob"},
    "summaries": {
        "0": "Diskussion:\nDie Sitzung wurde eröffnet.\nBeschluss:\nKeine Beschlüsse.",
        "1": (
            "Diskussion:\nDer Haushalt wurde beraten.\n"
            "Beschluss:\nDer Haushalt wurde empfohlen.\n"
            "Abstimmung:\nEinstimmig.\n"
            "Maßnahmen:\nDie Verwaltung legt Zahlen nach."
        ),
    },
    "summary_reviews": {},
}


def test_parse_summary_sections_extracts_professional_minutes_fields():
    sections = parse_summary_sections(EXPORT_PAYLOAD["summaries"]["1"])

    assert sections["discussion"] == ["Der Haushalt wurde beraten."]
    assert sections["decisions"] == ["Der Haushalt wurde empfohlen."]
    assert sections["votes"] == ["Einstimmig."]
    assert sections["action_items"] == ["Die Verwaltung legt Zahlen nach."]


def test_txt_export_contains_metadata_agenda_sections_and_appendix():
    with TestClient(main.app) as client:
        response = client.post("/api/export", json=EXPORT_PAYLOAD)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/plain")
    content = response.text
    assert "Gremium: Hauptausschuss" in content
    assert "Teilnehmer: Alice Beispiel, Bob Beispiel" in content
    assert "Begrüßung" in content
    assert "1. Begrüßung" not in content
    assert "Haushalt" in content
    assert "Beschluss:" in content
    assert "- Der Haushalt wurde empfohlen." in content
    assert "Abstimmung:" in content
    assert "Maßnahmen:" in content
    assert "Sprecherliste:" in content
    assert "Transkript:" in content
    assert "Ich eröffne die Sitzung." in content
    assert "Der Haushalt wird beraten." in content
    assert "Bearbeitungs-/Generierungshinweis:" in content


def test_generated_actions_and_open_questions_keep_categories_through_export():
    from summarize import StructuredSummary, render_structured_summary
    from export_protocol import build_protocol_document, ProtocolMetadata, render_protocol
    structured = StructuredSummary(action_items=['Die Verwaltung prüft den Antrag.'],
                                   open_points=['Die Finanzierung bleibt offen.'])
    text = render_structured_summary(structured)
    sections = parse_summary_sections(text)
    assert sections['action_items'] == structured.action_items
    assert sections['open_points'] == structured.open_points
    document = build_protocol_document(metadata=ProtocolMetadata(), tops=['Beratung'], summaries={0:text})
    txt = render_protocol(document, 'txt').decode()
    docx = Document(BytesIO(render_protocol(document, 'docx')))
    docx_text = '\n'.join(p.text for p in docx.paragraphs)
    import pdfplumber
    with pdfplumber.open(BytesIO(render_protocol(document, 'pdf'))) as pdf:
        pdf_text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
    for exported in [txt, docx_text, pdf_text]:
        assert 'Maßnahmen' in exported and 'Offene Punkte' in exported
        assert 'Finanzierung bleibt offen' in exported
        assert 'Maßnahmen/offene Punkte' not in exported


def test_legacy_combined_heading_is_preserved_without_guessing_categories():
    from export_protocol import build_protocol_document, ProtocolMetadata, render_protocol
    document = build_protocol_document(metadata=ProtocolMetadata(), tops=['Beratung'],
        summaries={0:'Maßnahmen/offene Punkte:\nDie Finanzierung bleibt offen.'})
    assert 'Maßnahmen/offene Punkte:' in render_protocol(document, 'txt').decode()


def test_txt_export_can_group_full_transcript_by_top():
    payload = {
        **EXPORT_PAYLOAD,
        "appendix": {
            **EXPORT_PAYLOAD["appendix"],
            "group_transcript_by_top": True,
        },
    }

    with TestClient(main.app) as client:
        response = client.post("/api/export", json=payload)

    assert response.status_code == 200
    content = response.text
    assert "Transkript:" in content
    assert "Begrüßung" in content
    assert "Haushalt" in content
    assert "TOP: TOP" not in content
    assert "[0:00] Alice: Ich eröffne die Sitzung." in content
    assert "[0:03] Bob: Der Haushalt wird beraten." in content


def test_legacy_transcript_excerpt_flag_exports_full_transcript():
    payload = {
        **EXPORT_PAYLOAD,
        "appendix": {
            "include_speaker_list": False,
            "include_transcript_excerpt": True,
            "transcript_excerpt_limit": 1,
            "include_generation_note": False,
        },
    }

    with TestClient(main.app) as client:
        response = client.post("/api/export", json=payload)

    assert response.status_code == 200
    content = response.text
    assert "Transkript:" in content
    assert "Ich eröffne die Sitzung." in content
    assert "Der Haushalt wird beraten." in content


def test_docx_export_has_expected_document_structure():
    payload = {**EXPORT_PAYLOAD, "format": "docx"}

    with TestClient(main.app) as client:
        response = client.post("/api/export", json=payload)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    doc = Document(BytesIO(response.content))
    paragraphs = [paragraph.text for paragraph in doc.paragraphs]

    assert "Sitzung Hauptausschuss" in paragraphs
    assert "Tagesordnung" in paragraphs
    assert "Begrüßung" in paragraphs
    assert "1. Begrüßung" not in paragraphs
    assert "Haushalt" in paragraphs
    assert "Beschluss" in paragraphs
    assert "Der Haushalt wurde empfohlen." in paragraphs
    assert "Anhang" in paragraphs

    table_text = "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    assert "Gremium" in table_text
    assert "Hauptausschuss" in table_text
    assert "Alice Beispiel, Bob Beispiel" in table_text


def test_pdf_export_returns_pdf_bytes_with_attachment_header():
    payload = {**EXPORT_PAYLOAD, "format": "pdf"}

    with TestClient(main.app) as client:
        response = client.post("/api/export", json=payload)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/pdf")
    assert response.headers["content-disposition"].endswith('.pdf"')
    assert response.content.startswith(b"%PDF")


def test_export_rejects_unknown_format():
    payload = {**EXPORT_PAYLOAD, "format": "xlsx"}

    with TestClient(main.app) as client:
        response = client.post("/api/export", json=payload)

    assert response.status_code == 400


def test_export_preserves_original_numbers_sections_and_unnumbered_legacy_titles():
    tops = ['[Öffentlich] 2.1 Schulbau', '[Nichtöffentlich] 2.1 Vergabe', '7 Anfragen', 'Alttitel']
    payload = {**EXPORT_PAYLOAD, 'tops': tops, 'summaries': {str(i): 'Diskussion:\nBeraten.' for i in range(4)}}
    client = TestClient(main.app)
    text = client.post('/api/export', json=payload).text
    assert all(top in text for top in tops)
    assert 'TOP 1:' not in text
    assert '1. [Öffentlich]' not in text
    response = client.post('/api/export', json={**payload, 'format': 'docx'})
    paragraphs = [paragraph.text for paragraph in Document(BytesIO(response.content)).paragraphs]
    assert all(top in paragraphs for top in tops)
