from io import BytesIO

from docx import Document
import pdfplumber
import pytest

from export_protocol import ProtocolAppendix, ProtocolMetadata, ProtocolTop, ProtocolDocument, render_protocol
from protocol_layout import parse_protocol, plain, render_text


EXAMPLE = '''## Zu TOP 02: Entscheidung über eventuelle Einwendungen gegen die Niederschrift

Der Ausschuss beschließt einstimmig (9 : 0 : 0) die Niederschrift vom 19.01.2026.

Aus der Beratung

Heide (SPD) bittet, die Ausführungen zu TOP 7 zur Kleinen Elster nochmals zu prüfen.

Frau Stasch werde die entsprechende Stelle prüfen.

Abstimmung:

| Abstimmungsergebnis | Abstimmungsergebnis | Abstimmungsergebnis | Abstimmungsergebnis |
|---------------------|---------------------|---------------------|---------------------|
|                     | JA                  | NEIN                | ENTH.               |
| AUK                 | 9                   | 0                   | 0                   |
'''


def document(text=EXAMPLE):
    return ProtocolDocument(metadata=ProtocolMetadata(title='Sitzung'), agenda=['TOP 02'], speakers=[], transcript=[],
        assignments=[], appendix=ProtocolAppendix(False, False, False, False), tops=[ProtocolTop(
            index=1, title='TOP 02', protocol_text=text, review_status='Ungeprüfter Entwurf.',
            review_questions=['Bitte die Angabe zu TOP 7 prüfen.'])])


@pytest.mark.parametrize('format', ['txt', 'docx', 'pdf'])
def test_protocol_exports_format_headings_and_vote_table_without_changing_source(format):
    source = document()
    artifact = render_protocol(source, format)
    assert source.tops[0].protocol_text == EXAMPLE
    if format == 'txt':
        text = artifact.decode('utf-8')
        assert 'AUK' in text and 'JA' in text and 'NEIN' in text and 'ENTH.' in text
        assert ' | ' in text and '-+-' in text
    elif format == 'docx':
        doc = Document(BytesIO(artifact))
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
        assert any(h.startswith('Zu TOP 02:') for h in headings)
        assert 'Aus der Beratung' in headings and 'Abstimmung:' in headings
        table = doc.tables[-1]
        assert [[cell.text for cell in row.cells] for row in table.rows] == [
            ['Abstimmungsergebnis'] * 4, ['', 'JA', 'NEIN', 'ENTH.'], ['AUK', '9', '0', '0']]
        assert table.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
        text = '\n'.join(p.text for p in doc.paragraphs)
    else:
        with pdfplumber.open(BytesIO(artifact)) as pdf:
            text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
            tables = [table for page in pdf.pages for table in page.extract_tables()]
            assert any(table[-1] == ['AUK', '9', '0', '0'] for table in tables)
    assert '## ' not in text and '|---' not in text
    assert 'Heide (SPD)' in text and 'Kleinen Elster' in text and 'Frau Stasch' in text
    assert '(9 : 0 : 0)' in text and '19.01.2026' in text
    assert 'Bitte die Angabe zu TOP 7 prüfen.' in text
    assert 'Keine Angabe.' not in text


def test_inline_emphasis_numbering_links_and_escaped_pipe_cells_are_retained():
    value = '**Aus der Beratung**\n\nEin **wichtiger** *Hinweis* und ~~veraltet~~.\n\n3. Erster Punkt\n4. Zweiter Punkt\n\n| Name | Zahl |\n|---|---|\n| A \\| B | 12 |\n\n[Plan](https://example.org/plan)\n\nLetzter Absatz.'
    blocks = parse_protocol(value)
    assert blocks[0].kind == 'heading'
    assert any(run.bold and run.text == 'wichtiger' for block in blocks for run in block.runs)
    assert any(run.italic and run.text == 'Hinweis' for block in blocks for run in block.runs)
    assert any(run.strike and run.text == 'veraltet' for block in blocks for run in block.runs)
    assert [block.prefix for block in blocks if block.prefix] == ['3. ', '4. ']
    table = next(block for block in blocks if block.kind == 'table')
    assert [plain(cell) for cell in table.rows[1]] == ['A | B', '12']
    text = render_text(value)
    assert 'https://example.org/plan' in text and text.endswith('Letzter Absatz.')
    doc = Document(BytesIO(render_protocol(document(value), 'docx')))
    assert any(run.bold and run.text == 'wichtiger' for paragraph in doc.paragraphs for run in paragraph.runs)
    assert any(run.italic and run.text == 'Hinweis' for paragraph in doc.paragraphs for run in paragraph.runs)
    assert any(run.font.strike and run.text == 'veraltet' for paragraph in doc.paragraphs for run in paragraph.runs)
    render_protocol(document(value), 'pdf')


def test_malformed_markdown_and_html_are_preserved_as_readable_text():
    value = 'Erster Absatz.\nZweite Zeile.\n\n| A | B |\nKein Tabellentrenner.\n\n<script>alert(1)</script>\n\nLetzter Absatz.'
    assert not any(block.kind == 'table' for block in parse_protocol(value))
    text = render_text(value)
    assert '| A | B |' in text and '<script>alert(1)</script>' in text
    assert text.endswith('Letzter Absatz.')
    render_protocol(document(value), 'pdf')  # ReportLab receives escaped text, not model-generated XML.


def test_long_tables_span_pdf_pages_and_keep_every_final_row():
    table = '| Auftrag | Zahl |\n|---|---|\n' + '\n'.join(f'| Auftrag {i} | {i} |' for i in range(100))
    with pdfplumber.open(BytesIO(render_protocol(document(table), 'pdf'))) as pdf:
        assert len(pdf.pages) > 1
        text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
        for i in range(100):
            assert f'Auftrag {i}' in text
        assert all('Auftrag Zahl' in (page.extract_text() or '') for page in pdf.pages[1:-1])
        for page in pdf.pages:
            content = page.extract_text() or ''
            if 'Prüfhinweise' in content:
                assert 'Bitte die Angabe zu TOP 7 prüfen.' in content


def test_large_table_cell_splits_across_pages_without_losing_its_end():
    import re
    value = '| Text | Zahl |\n|---|---|\n| Anfang ' + 'Langer Text. ' * 500 + 'Ende | 9 |'
    with pdfplumber.open(BytesIO(render_protocol(document(value), 'pdf'))) as pdf:
        assert len(pdf.pages) > 1
        text = re.sub(r'\s+', ' ', ' '.join(page.extract_text() or '' for page in pdf.pages))
        # A repeated header/footer can fall between words of a split cell.
        assert text.count('Langer') == 500 and text.count('Text.') == 500
        assert 'Anfang' in text and 'Ende' in text
